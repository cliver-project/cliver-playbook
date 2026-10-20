"""Tests for the DOCX chunker."""

from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document as DocxDocument
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from cliver_playbook.app.rag.chunking import Document, get_chunker
from cliver_playbook.app.rag.chunking.base import ChunkerInputError, EmptyDocumentError
from cliver_playbook.app.rag.chunking.options import DocxOptions


def _build(builder) -> bytes:
    document = DocxDocument()
    builder(document)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _split(raw: bytes, **kwargs):
    doc = Document(
        id="docx-1",
        source_version="v1",
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        raw=raw,
    )
    return get_chunker("docx").split(doc, **kwargs)


def _texts(chunks) -> str:
    return "\n".join(chunk.text for chunk in chunks)


def test_headings_become_sections_with_paths() -> None:
    def builder(document):
        document.add_heading("Policy", level=1)
        document.add_paragraph("intro")
        document.add_heading("Scope", level=2)
        document.add_paragraph("scope body")

    chunks = _split(_build(builder), chunk_size=512, chunk_overlap=0)

    assert [chunk.metadata["section_path"] for chunk in chunks] == [
        ["H1#Policy"],
        ["H1#Policy", "H2#Scope"],
    ]
    assert chunks[0].text == "Policy\n\nintro"
    assert chunks[1].text == "Policy > Scope\n\nscope body"


def test_outline_level_is_honoured_before_style_name() -> None:
    def builder(document):
        paragraph = document.add_paragraph("Outline heading")
        p_pr = paragraph._p.get_or_add_pPr()
        outline = OxmlElement("w:outlineLvl")
        outline.set(qn("w:val"), "2")
        p_pr.append(outline)
        document.add_paragraph("body")

    chunks = _split(_build(builder), chunk_size=512, chunk_overlap=0)

    assert chunks[0].metadata["section_path"] == ["H3#Outline heading"]
    assert chunks[0].text == "Outline heading\n\nbody"


def test_toc_paragraphs_are_dropped_by_default() -> None:
    def builder(document):
        style = document.styles.add_style("TOC 1", WD_STYLE_TYPE.PARAGRAPH)
        document.add_paragraph("Chapter 1 .... 1", style=style)
        document.add_heading("Real", level=1)
        document.add_paragraph("body")

    raw = _build(builder)
    chunks = _split(raw, chunk_size=512, chunk_overlap=0)
    assert "Chapter 1" not in _texts(chunks)

    kept = _split(raw, chunk_size=512, chunk_overlap=0, options=DocxOptions(drop_toc=False))
    assert "Chapter 1" in _texts(kept)


def test_tables_render_as_markdown_and_html() -> None:
    def builder(document):
        document.add_heading("Data", level=1)
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "name"
        table.cell(0, 1).text = "value"
        table.cell(1, 0).text = "alpha"
        table.cell(1, 1).text = "1"

    raw = _build(builder)
    markdown = _texts(_split(raw, chunk_size=512, chunk_overlap=0))
    html = _texts(
        _split(
            raw,
            chunk_size=512,
            chunk_overlap=0,
            options=DocxOptions(table_format="html"),
        )
    )

    assert "| name | value |" in markdown
    assert "| alpha | 1 |" in markdown
    assert "<table>" in html
    assert "<td>alpha</td>" in html


def test_tables_can_be_excluded() -> None:
    def builder(document):
        document.add_paragraph("before")
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "cell"
        document.add_paragraph("after")

    chunks = _split(
        _build(builder),
        chunk_size=512,
        chunk_overlap=0,
        options=DocxOptions(include_tables=False),
    )

    assert _texts(chunks) == "before\n\nafter"


def test_document_order_is_preserved() -> None:
    def builder(document):
        document.add_paragraph("before")
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "cell"
        document.add_paragraph("after")

    text = _texts(_split(_build(builder), chunk_size=512, chunk_overlap=0))

    assert text.index("before") < text.index("cell") < text.index("after")


def test_heading_based_can_be_disabled() -> None:
    def builder(document):
        document.add_heading("Policy", level=1)
        document.add_paragraph("body")

    chunks = _split(
        _build(builder),
        chunk_size=512,
        chunk_overlap=0,
        options=DocxOptions(heading_based=False),
    )

    assert len(chunks) == 1
    assert chunks[0].text == "Policy\n\nbody"
    assert chunks[0].metadata["section_path"] == []


def test_max_heading_level_turns_deep_headings_into_text() -> None:
    def builder(document):
        document.add_heading("Policy", level=1)
        document.add_heading("Scope", level=2)
        document.add_paragraph("body")

    chunks = _split(
        _build(builder),
        chunk_size=512,
        chunk_overlap=0,
        options=DocxOptions(max_heading_level=1),
    )

    assert chunks[0].metadata["section_path"] == ["H1#Policy"]
    assert chunks[0].text == "Policy\n\nScope\n\nbody"


def test_heading_prefix_can_be_disabled() -> None:
    def builder(document):
        document.add_heading("Policy", level=1)
        document.add_paragraph("body")

    chunks = _split(
        _build(builder),
        chunk_size=512,
        chunk_overlap=0,
        options=DocxOptions(include_heading_path=False),
    )

    assert chunks[0].text == "body"
    assert chunks[0].metadata["section_path"] == ["H1#Policy"]


def test_table_only_document_still_produces_chunks() -> None:
    def builder(document):
        table = document.add_table(rows=2, cols=1)
        table.cell(0, 0).text = "header"
        table.cell(1, 0).text = "value"

    chunks = _split(_build(builder), chunk_size=512, chunk_overlap=0)

    assert "| header |" in chunks[0].text
    assert "| value |" in chunks[0].text


def test_legacy_doc_is_rejected() -> None:
    raw = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"legacy"

    with pytest.raises(ChunkerInputError, match="\\.docx"):
        _split(raw)


def test_non_docx_bytes_are_rejected() -> None:
    with pytest.raises(ChunkerInputError):
        _split(b"this is not a zip package")


def test_empty_docx_is_rejected() -> None:
    with pytest.raises(EmptyDocumentError):
        _split(_build(lambda document: None))
